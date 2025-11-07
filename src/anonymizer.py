
# ==== הגדרות ====
import os, re, subprocess, shlex, hashlib
from pathlib import Path
import pytesseract
pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

# <<< נתיבי קלט/פלט >>>  ---- עדכני לפי הצורך ----
INPUT_DIR  = '/content/drive/MyDrive/CV/all'
OUTPUT_DIR = '/content/drive/MyDrive/CV_all_fina1_ocr'
Path(OUTPUT_DIR).mkdir(parents=True, exist_ok=True)

# ===== תבניות ערכים =====
PHONE_VAL_TXT = r'(?:\+?972[-\s]?|0)(?:[23489][- \s]?\d{7}|5\d[- \s]?\d{3}[- \s]?\d{4})'
ID_VAL_TXT    = r'(?<!\d)\d{7,10}(?!\d)'
EMAIL_VAL_TXT = r'[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}'
# קישורים: http/https, mailto, www, וגם דומיין חשוף (example.co.il/…)
URL_VAL_TXT   = r'(?:(?:https?://|mailto:|www\.)[^\s<>()]+|[A-Za-z0-9.-]+\.[A-Za-z]{2,}[^\s<>()]*)'

# ===== תוויות =====
LABEL_PHONE = r'(?:נייד|טלפון|Phone|Cell|Mobile)'
LABEL_ID    = r'(?:ת\.?\s*ז|תז|תעודת\s*זהות|מספר\s*זהות|ID)'
LABEL_EMAIL = r'(?:מייל|אימייל|דוא"?ל|דואר\s*אלקטרוני|Email|E-?mail)'
LABEL_NAME  = r'(?:שם(?:\s*מלא)?|Name|Full\s*Name)'
LABEL_FNAME = r'(?:שם\s*פרטי|First\s*Name)'
LABEL_LNAME = r'(?:שם\s*משפחה|Last\s*Name)'

# --- Birthdate -> Year ---
LABEL_BIRTHDATE = r'(?:תאריך\s*לידה|ת\.?\s*לידה|תאריך\s*הולדת|DOB|Date\s*of\s*Birth|Birth\s*Date|Birthdate|Born)'
HEB_MONTHS = r'(?:ינואר|פברואר|מרץ|אפריל|מאי|יוני|יולי|אוגוסט|ספטמבר|אוקטובר|נובמבר|דצמבר)'
ENG_MONTHS = r'(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|Jul(?:y)?|Aug(?:ust)?|Sep(?:t(?:ember)?)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)'

REPL = '[Confidential]'
WSP  = r'[\s\u00A0\u200E\u200F]*'  # רווח/NBSP/RTL/LTR

CV_TITLE_RE = re.compile(r'קורות' + WSP + r'[-־–—]?' + WSP + r'חיים')

SECTION_WORDS = {
    'פרטים','אישיים','ניסיון','השכלה','מיומנויות','כישורים','שפות',
    'המלצות','קורסים','תעסוקה','תעסוקתי','פרויקטים','תקציר',
    'סיכום','מטרה','הסמכות','תעודות'
}

def _clean(s: str) -> str:
    if not s: return ''
    return (s.replace('\xa0',' ')
             .replace('\u200f','').replace('\u200e','').replace('\u200b','')
             .strip())

def _is_probably_he_name(s: str) -> bool:
    t = re.sub(r'[\"“”\'׳״.,|•:;()\[\]\-–—]', ' ', _clean(s))
    t = re.sub(r'\s+', ' ', t).strip()
    if not t or any(ch.isdigit() for ch in t) or len(t) > 40:
        return False
    words = [w for w in t.split() if re.fullmatch(r'[א-ת]+', w)]
    return 2 <= len(words) <= 4 and not any(w in SECTION_WORDS for w in words)

# ===== DOCX =====
from docx import Document
from docx.oxml.ns import qn

def _set_para_text(para, text):
    for r in para.runs: r.text = ''
    para.add_run(text)

def _normalize_title_or_name(text: str) -> str:
    s = _clean(text)
    if not s: return text
    if CV_TITLE_RE.search(s): return 'קורות חיים'
    if _is_probably_he_name(s): return REPL
    return text

def _yy_to_yyyy(yy: str, cutoff: int = 24) -> str:
    n = int(yy)
    return f'20{n:02d}' if n <= cutoff else f'19{n:02d}'

def _sanitize_cv_title(text: str) -> str:
    s = _clean(text)
    m = CV_TITLE_RE.search(s)
    if not m: return text
    extras = (s[:m.start()] + s[m.end():]).strip(' :|-–—•.,\u00a0')
    return 'קורות חיים' if extras else CV_TITLE_RE.sub('קורות חיים', s)

def _mask_name_inline(s: str) -> str:
    t = s
    for lab, can in ((LABEL_FNAME, 'שם פרטי'),
                     (LABEL_LNAME, 'שם משפחה'),
                     (LABEL_NAME,  'שם')):
        t2 = re.sub(rf'(?<!\S){lab}{WSP}[:：\-]?{WSP}.+', f'{can}: {REPL}', t, flags=re.IGNORECASE)
        if t2 != t: t = t2
    return t

def _birthdate_to_year(s: str) -> str:
    t = s
    # עם תווית
    t = re.sub(rf'({LABEL_BIRTHDATE}){WSP}[:：\-]?{WSP}\d{{1,2}}[./\-]\d{{1,2}}[./\-](\d{{4}})', rf'\1: \2', t, flags=re.IGNORECASE)
    t = re.sub(rf'({LABEL_BIRTHDATE}){WSP}[:：\-]?{WSP}(\d{{4}})[./\-]\d{{1,2}}[./\-]\d{{1,2}}', rf'\1: \2', t, flags=re.IGNORECASE)
    t = re.sub(rf'({LABEL_BIRTHDATE}){WSP}[:：\-]?{WSP}\d{{1,2}}{WSP}({HEB_MONTHS}){WSP}(\d{{4}})', rf'\1: \3', t, flags=re.IGNORECASE)
    t = re.sub(rf'({LABEL_BIRTHDATE}){WSP}[:：\-]?{WSP}{ENG_MONTHS}{WSP}\d{{1,2}}(?:,)?{WSP}(\ד{{4}})'.replace('ד','\\d'), rf'\1: \2', t, flags=re.IGNORECASE)
    # YY
    t = re.sub(rf'({LABEL_BIRTHDATE}){WSP}[:：\-]?{WSP}\ד{{1,2}}[./\-]\ד{{1,2}}[./\-](\ד{{2}})'.replace('ד','\\d'),
               lambda m: f"{m.group(1)}: {_yy_to_yyyy(m.group(2))}", t, flags=re.IGNORECASE)
    t = re.sub(rf'({LABEL_BIRTHDATE}){WSP}[:：\-]?{WSP}{ENG_MONTHS}{WSP}\ד{{1,2}}(?:,)?{WSP}(\ד{{2}})'.replace('ד','\\d'),
               lambda m: f"{m.group(1)}: {_yy_to_yyyy(m.group(2))}", t, flags=re.IGNORECASE)
    # בלי תווית
    t = re.sub(rf'(נולד(?:ה)?){WSP}\ד{{1,2}}[./\-]\ד{{1,2}}[./\-](\ד{{4}})'.replace('ד','\\d'), rf'\1 \2', t)
    t = re.sub(rf'(Born){WSP}(?:on{WSP})?(\ד{{4}})[./\-]\ד{{1,2}}[./\-]\ד{{1,2}}'.replace('ד','\\d'), rf'\1 \2', t, flags=re.IGNORECASE)
    # YY בלי תווית
    t = re.sub(rf'(נולד(?:ה)?){WSP}\ד{{1,2}}[./\-]\ד{{1,2}}[./\-](\ד{{2}})'.replace('ד','\\d'), lambda m: f"{m.group(1)} {_yy_to_yyyy(m.group(2))}", t)
    t = re.sub(rf'(Born){WSP}(?:on{WSP})?(\ד{{2}})[./\-]\ד{{1,2}}[./\-]\ד{{1,2}}'.replace('ד','\\d'), lambda m: f"{m.group(1)} {_yy_to_yyyy(m.group(2))}", t, flags=re.IGNORECASE)
    return t

def replace_in_line(line: str) -> str:
    s = _clean(line)
    new_title = _sanitize_cv_title(s)
    if new_title != s: return new_title
    s2 = _mask_name_inline(s)
    if s2 != s: s = s2
    # תווית + ערך
    s = re.sub(rf'({LABEL_PHONE}){WSP}[:：\-]?{WSP}{PHONE_VAL_TXT}',  rf'\1: {REPL}', s, flags=re.IGNORECASE)
    s = re.sub(rf'({LABEL_ID}){WSP}[:：\-]?{WSP}{ID_VAL_TXT}',      rf'\1: {REPL}', s, flags=re.IGNORECASE)
    s = re.sub(rf'({LABEL_EMAIL}){WSP}[:：\-]?{WSP}{EMAIL_VAL_TXT}', rf'\1: {REPL}', s, flags=re.IGNORECASE)
    # ערכים בודדים + קישורים
    s = re.sub(rf'mailto:{EMAIL_VAL_TXT}', REPL, s, flags=re.IGNORECASE)
    s = re.sub(EMAIL_VAL_TXT, REPL, s, flags=re.IGNORECASE)
    s = re.sub(URL_VAL_TXT, REPL, s, flags=re.IGNORECASE)
    s = re.sub(PHONE_VAL_TXT, REPL, s)
    s = re.sub(ID_VAL_TXT,     REPL, s)
    s = _birthdate_to_year(s)
    return s

def merge_label_value_paragraphs(paragraphs):
    LABEL_ONLY = {
        'נייד':       re.compile(rf'^{WSP}({LABEL_PHONE}){WSP}[:：\-]?{WSP}$', re.IGNORECASE),
        'תז':         re.compile(rf'^{WSP}({LABEL_ID}){WSP}[:：\-]?{WSP}$',    re.IGNORECASE),
        'מייל':       re.compile(rf'^{WSP}({LABEL_EMAIL}){WSP}[:：\-]?{WSP}$', re.IGNORECASE),
        'שם':         re.compile(rf'^{WSP}({LABEL_NAME}){WSP}[:：\-]?{WSP}$',  re.IGNORECASE),
        'שם פרטי':    re.compile(rf'^{WSP}({LABEL_FNAME}){WSP}[:：\-]?{WSP}$', re.IGNORECASE),
        'שם משפחה':   re.compile(rf'^{WSP}({LABEL_LNAME}){WSP}[:：\-]?{WSP}$', re.IGNORECASE),
        'תאריך לידה': re.compile(rf'^{WSP}({LABEL_BIRTHDATE}){WSP}[:：\-]?{WSP}$', re.IGNORECASE),
    }
    COLON_ONLY = re.compile(rf'^{WSP}[:：\-]{WSP}$')

    i = 0
    while i < len(paragraphs):
        t = _clean(paragraphs[i].text)
        if i < 6:
            new_t = _normalize_title_or_name(t)
            if new_t != t: _set_para_text(paragraphs[i], new_t)
            t = _clean(paragraphs[i].text)
        if not t: i += 1; continue

        t_inline = _mask_name_inline(t)
        if t_inline != t:
            _set_para_text(paragraphs[i], t_inline); i += 1; continue

        matched_key = None
        for key, rgx in LABEL_ONLY.items():
            if rgx.match(t): matched_key = key; break
        if matched_key:
            j = i + 1
            while j < len(paragraphs):
                tj = _clean(paragraphs[j].text)
                if not tj or COLON_ONLY.match(tj): j += 1; continue
                if matched_key == 'תאריך לידה':
                    year_line = _birthdate_to_year(tj)
                    out_val = year_line if year_line != tj else 'תאריך לידה: '
                    _set_para_text(paragraphs[i], out_val)
                else:
                    _set_para_text(paragraphs[i], f'{matched_key}: {REPL}')
                for k in range(i+1, j+1): _set_para_text(paragraphs[k], '')
                break
            i = j; continue

        new = replace_in_line(t)
        if new != t: _set_para_text(paragraphs[i], new)
        i += 1

def _safe_drop_rel(part, rid):
    try:
        part.drop_rel(rid)
    except Exception:
        try:
            if hasattr(part, '_rels') and rid in part._rels:
                del part._rels[rid]
        except Exception:
            pass

def _scrub_hyperlinks_in_part(part):
    if not hasattr(part, 'element') or not hasattr(part, 'rels'):
        return
    rels = part.rels
    for h in part.element.xpath('.//w:hyperlink'):
        rid = h.get(qn('r:id'))
        target = ''
        if rid and rid in rels:
            target = getattr(rels[rid], 'target_ref', '') or getattr(rels[rid], '_target', '')
        display_text = ''.join(t.text or '' for t in h.xpath('.//w:t'))
        need_scrub = False
        if re.search(EMAIL_VAL_TXT, display_text or '', re.IGNORECASE) or re.search(URL_VAL_TXT, display_text or '', re.IGNORECASE):
            need_scrub = True
        if target and (target.lower().startswith('mailto:') or
                       re.search(EMAIL_VAL_TXT, target, re.IGNORECASE) or
                       re.search(URL_VAL_TXT, target, re.IGNORECASE)):
            need_scrub = True
        if need_scrub:
            for t in h.xpath('.//w:t'): t.text = REPL
            if rid:
                _safe_drop_rel(part, rid)
            try: h.attrib.pop(qn('r:id'), None)
            except Exception: pass

def anonymize_docx_file(path_in: str, path_out: str):
    doc = Document(path_in)
    # נטרול היפרלינקים
    _scrub_hyperlinks_in_part(doc.part)
    for sec in doc.sections:
        if hasattr(sec, 'header') and hasattr(sec.header, 'part'):
            _scrub_hyperlinks_in_part(sec.header.part)
        if hasattr(sec, 'footer') and hasattr(sec.footer, 'part'):
            _scrub_hyperlinks_in_part(sec.footer.part)

    # headers/footers
    for sec in doc.sections:
        for para in list(sec.header.paragraphs): _process_paragraph_obj(para)
        for tbl in list(sec.header.tables):      _process_table_obj(tbl)
        for para in list(sec.footer.paragraphs): _process_paragraph_obj(para)
        for tbl in list(sec.footer.tables):      _process_table_obj(tbl)

    # גוף + טבלאות
    merge_label_value_paragraphs(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                before = _clean(cell.text)
                if not before: continue
                fixed = _normalize_title_or_name(before)
                if fixed != before:
                    cell.text = fixed; continue
                after = replace_in_line(before)
                after = _birthdate_to_year(after)
                if after != before: cell.text = after
                if _clean(cell.text) in {':','-','–','־'}: cell.text = ''

    doc.save(path_out)

def _process_paragraph_obj(para):
    before = _clean(para.text)
    if not before: return
    fixed = _normalize_title_or_name(before)
    if fixed != before:
        _set_para_text(para, fixed); return
    after = replace_in_line(before)
    after = _birthdate_to_year(after)
    if after != before: _set_para_text(para, after)

def _process_table_obj(table):
    for row in table.rows:
        for cell in row.cells:
            t0 = _clean(cell.text)
            if not t0: continue
            t1 = _normalize_title_or_name(t0)
            if t1 != t0:
                cell.text = t1
            else:
                t2 = replace_in_line(t0)
                t2 = _birthdate_to_year(t2)
                if t2 != t0: cell.text = t2

# ===== PDF (PyMuPDF) =====
import fitz

PHONE_VAL_PDF = r'(?:\+?972[-\s\u2010-\u2014]?|0)(?:[23489][- \s\u2010-\u2014]?\d{7}|5\d[- \s\u2010-\u2014]?\d{3}[- \s\u2010-\u2014]?\d{4})'
ID_VAL_PDF    = r'(?<!\d)\d{7,10}(?!\d)'
EMAIL_VAL_PDF = EMAIL_VAL_TXT
URL_VAL_PDF   = URL_VAL_TXT

# --- מיזוג מלבנים לפי שורה (מונע "מחיקה של כל הדף") ---
def _merge_overlapping_rects_linewise(rects, y_tol=6):
    """ממזג רק מלבנים שחופפים משמעותית באותו קו טקסט (y-mid קרוב).
       לא ממזג שורות שונות כדי למנוע איחוד-על."""
    rects = [fitz.Rect(r) for r in rects]
    if not rects: return []
    # מסדרים לפי מרכז Y
    def y_mid(r): return (r.y0 + r.y1)/2
    buckets = []
    for r in sorted(rects, key=y_mid):
        placed = False
        for b in buckets:
            if abs(y_mid(b[0]) - y_mid(r)) <= y_tol:
                b.append(r); placed = True; break
        if not placed: buckets.append([r])
    merged = []
    for bucket in buckets:
        bucket.sort(key=lambda rr: (rr.x0, rr.x1))
        out = []
        for r in bucket:
            if not out:
                out.append(r); continue
            last = out[-1]
            # מאחד אם יש חפיפה/מגע קטן אופקי ושייכות לשורה
            if r.x0 <= last.x1 + 2 and r.y1 >= last.y0 - y_tol and r.y0 <= last.y1 + y_tol:
                out[-1] = last | r
            else:
                out.append(r)
        merged.extend(out)
    return merged

def _words(page):
    ws = page.get_text("words")
    ws.sort(key=lambda w: (w[5], w[6], w[7]))
    return ws

def _rect_union(words, i_from, i_to):
    r = fitz.Rect(words[i_from][0], words[i_from][1], words[i_from][2], words[i_from][3])
    for k in range(i_from+1, i_to+1): r |= fitz.Rect(words[k][0], words[k][1], words[k][2], words[k][3])
    return r

# שם כשדה ב-PDF (בטקסט אמיתי)
LABEL_TOKENS_SINGLE = {'שם'}
LABEL_TOKENS_PAIR   = {('שם','מלא'), ('שם','פרטי'), ('שם','משפחה')}
COLON_TOKENS = {':','־','-','–','—'}
NEXT_LABEL_STARTERS = {'תז','נייד','טלפון','מייל','אימייל','Email','Phone','תאריך','כתובת'}

def _redact_name_values_rects(page):
    ws = _words(page); rects = []; n = len(ws); i = 0
    def same_block(j, w): return j < n and ws[j][5] == w[5]
    def same_line(j, w):  return j < n and ws[j][5] == w[5] and ws[j][6] == w[6]
    while i < n:
        w = ws[i]; tok = ws[i][4].strip()
        if tok in LABEL_TOKENS_SINGLE or tok.rstrip(':') in LABEL_TOKENS_SINGLE:
            j = i + 1
            while same_line(j, w) and ws[j][4].strip() in COLON_TOKENS: j += 1
            start = j
            if not same_line(start, w) and same_block(start, w):
                while same_block(start, w) and ws[start][4].strip() in COLON_TOKENS: start += 1
            if start < n and same_block(start, w):
                k = start
                while same_block(k, w) and ws[k][4].strip() not in NEXT_LABEL_STARTERS: k += 1
                rects.append(_rect_union(ws, start, k-1)); i = k; continue
        if i+1 < n and (tok, ws[i+1][4].strip()) in LABEL_TOKENS_PAIR and ws[i+1][5] == w[5]:
            j = i + 2
            while same_line(j, w) and ws[j][4].strip() in COLON_TOKENS: j += 1
            start = j
            if not same_line(start, w) and same_block(start, w):
                while same_block(start, w) and ws[start][4].strip() in COLON_TOKENS: start += 1
            if start < n and same_block(start, w):
                k = start
                while same_block(k, w) and ws[k][4].strip() not in NEXT_LABEL_STARTERS: k += 1
                rects.append(_rect_union(ws, start, k-1)); i = k; continue
        i += 1
    return rects

# === שם ככותרת (ללא תווית) בטקסט אמיתי ===
HE_NAME_LINE_RE = re.compile(
    r'^[\s\u200f\u200e]*([א-ת]+(?:\s+[א-ת]+){1,3})[\s\u200f\u200e]*$'
)
def _probable_he_name_line(s: str) -> bool:
    s = _clean(s)
    if not s or len(s) > 50: return False
    m = HE_NAME_LINE_RE.match(s)
    if not m: return False
    words = [w for w in re.sub(r'\s+', ' ', m.group(1)).split(' ') if w]
    bad = {'קורות','חיים','פרטים','אישיים','ניסיון','השכלה','מיומנויות','כישורים','שפות','תקציר','סיכום','מטרה','הסמכות','תעודות'}
    return 2 <= len(words) <= 4 and not any(w in bad for w in words)

def _redact_probable_name_title(page):
    blocks = page.get_text("blocks") or []
    if not blocks: return 0
    blocks.sort(key=lambda b: (b[1], b[0]))
    done = 0
    for (x0,y0,x1,y1,txt,*_) in blocks[:5]:
        if not txt: continue
        for line in [l.strip() for l in txt.split('\n') if _clean(l)]:
            if _probable_he_name_line(line):
                for r in page.search_for(line):
                    page.add_redact_annot(r, text=REPL, fill=(1,1,1)); done += 1
                if done: return done
    return done

# --- שם ככותרת לפי גודל פונטים (top of page) ---
def _redact_big_font_name_spans(page, top_ratio=0.35, min_font=18, max_words=4):
    info = page.get_text("dict") or {}
    page_h = page.rect.height
    hits = 0
    for block in info.get("blocks", []):
        for line in block.get("lines", []):
            for span in line.get("spans", []):
                text = _clean(span.get("text", ""))
                size = span.get("size", 0)
                bbox  = fitz.Rect(span.get("bbox", page.rect))
                if bbox.y1 > page_h * top_ratio:
                    continue
                if size >= min_font and re.fullmatch(r'[א-ת\s]{2,}', text or ''):
                    words = [w for w in re.sub(r'\s+', ' ', text).split(' ') if w]
                    if 2 <= len(words) <= max_words and not any(w in SECTION_WORDS for w in words):
                        page.add_redact_annot(bbox, text=REPL, fill=(1,1,1))
                        hits += 1
    return hits

# ===== OCR =====
import pytesseract
from PIL import Image
import io

def _ocr_words(page, zoom=3.0, lang='heb+eng'):
    mat = fitz.Matrix(zoom, zoom)
    pix = page.get_pixmap(matrix=mat, alpha=False)
    img = Image.open(io.BytesIO(pix.tobytes("png")))
    data = pytesseract.image_to_data(img, lang=lang, output_type=pytesseract.Output.DICT)
    words = []
    confs = data.get('conf', [])
    for i in range(len(data['text'])):
        txt = (data['text'][i] or '').strip()
        try:
            conf = int(confs[i])
        except Exception:
            conf = -1
        if not txt or conf < 0:
            continue
        x, y, w, h = data['left'][i], data['top'][i], data['width'][i], data['height'][i]
        x0, y0, x1, y1 = x/zoom, y/zoom, (x+w)/zoom, (y+h)/zoom
        words.append({'x0':x0, 'y0':y0, 'x1':x1, 'y1':y1, 'text':txt})
    words.sort(key=lambda w: (round(w['y0'], 2), w['x0']))
    return words

def _rect_union_ocr(ws, i_from, i_to):
    r = fitz.Rect(ws[i_from]['x0'], ws[i_from]['y0'], ws[i_from]['x1'], ws[i_from]['y1'])
    for k in range(i_from+1, i_to+1):
        r |= fitz.Rect(ws[k]['x0'], ws[k]['y0'], ws[k]['x1'], ws[k]['y1'])
    return r

def _ocr_find_token_windows(page, regex, max_tokens=8, zoom=3.0):
    ws = _ocr_words(page, zoom=zoom)
    rects = []
    n = len(ws)
    for i in range(n):
        for j in range(i, min(n, i+max_tokens)):
            s0 = ''.join(ws[k]['text'] for k in range(i, j+1))
            s1 = ' '.join(ws[k]['text'] for k in range(i, j+1))
            if regex.search(s0) or regex.search(s1):
                rects.append(_rect_union_ocr(ws, i, j)); break
    return rects

def _ocr_redact_probable_name_title(page, zoom=3.0):
    ws = _ocr_words(page, zoom=zoom)
    lines = {}
    for w in ws:
        key = round(w['y0'], 1)
        lines.setdefault(key, []).append(w)
    taken = 0
    for yk in sorted(lines.keys())[:3]:
        text_line = ' '.join([w['text'] for w in sorted(lines[yk], key=lambda z: z['x0'])])
        if _probable_he_name_line(text_line):
            rect = _rect_union_ocr(sorted(lines[yk], key=lambda z: z['x0']), 0, len(lines[yk])-1)
            page.add_redact_annot(rect, text=REPL, fill=(1,1,1)); taken += 1
            break
    return taken

def _ocr_birthdate_to_year_rects(page, zoom=3.0):
    ws = _ocr_words(page, zoom=zoom)
    tokens = [w['text'] for w in ws]
    rects = []
    date_patterns = [
        re.compile(r'\b(\d{4})[./\-]\d{1,2}[./\-]\d{1,2}\b'),
        re.compile(r'\b\d{1,2}[./\-]\d{1,2}[./\-](\d{4})\b'),
        re.compile(rf'\b{ENG_MONTHS}\s+\d{{1,2}},?\s+(\d{{4}})\b', re.IGNORECASE),
        re.compile(rf'\b\d{{1,2}}\s+{HEB_MONTHS}\s+(\d{{4}})\b'),
        re.compile(r'\b\d{1,2}[./\-]\d{1,2}[./\-](\d{2})\b'),
        re.compile(rf'\b{ENG_MONTHS}\s+\d{{1,2}},?\s+(\d{{2}})\b', re.IGNORECASE),
    ]
    label_re = re.compile(rf'({LABEL_BIRTHDATE})', re.IGNORECASE)
    n = len(tokens)
    for i in range(n):
        for j in range(i, min(n, i+8)):
            s0 = ''.join(tokens[i:j+1])
            s1 = ' '.join(tokens[i:j+1])
            date_m = None
            for pat in date_patterns:
                mm0 = pat.search(s0)
                mm1 = pat.search(s1) if not mm0 else None
                date_m = mm0 or mm1
                if date_m: break
            if not date_m:
                continue
            left_tokens  = ' '.join(tokens[max(0, i-6):i])
            right_tokens = ' '.join(tokens[j+1:min(n, j+7)])
            # גם בלי תווית נחליף לשנה
            year = None
            for gname in ('1','yyyy','yyyy2','yyyy3','yyyy4','yy','yy2'):
                try:
                    year = date_m.group(int(gname)) if gname=='1' else date_m.group(gname)
                    if year: break
                except Exception:
                    pass
            if not year: continue
            if len(year) == 2:
                year = f"20{int(year):02d}" if int(year) <= 24 else f"19{int(year):02d}"
            rect = _rect_union_ocr(ws, i, j)
            rects.append((rect, year))
            break
    return rects

# --- CAP בטיחות: לא נמרח על כל הדף ---
def _safe_apply_redactions(page, added_rects, cover_limit=0.70):
    """מונע מצב שמלבני המחיקה מכסים >70% מהעמוד"""
    if not added_rects:
        return
    # מאחדים לבאונדרי אחד (Bounding Box) – זה שמרני אבל בטוח
    union = None
    for r in added_rects:
        union = r if union is None else (union | r)
    # חישוב שטח בלי get_area (תואם לכל גרסאות PyMuPDF)
    page_area  = float(page.rect.width) * float(page.rect.height)
    union_area = float(union.width) * float(union.height) if union else 0.0
    # אם הבנאונדינג מכסה מעל הסף – אל תיישם מחיקות (מונע "עמוד ריק")
    if page_area > 0 and (union_area / page_area) > cover_limit:
        return
    page.apply_redactions()



def anonymize_pdf_file(path_in: str, path_out: str):
    doc = fitz.open(path_in)
    p_phone = re.compile(PHONE_VAL_PDF)
    p_id    = re.compile(ID_VAL_PDF)
    p_mail  = re.compile(EMAIL_VAL_PDF, re.IGNORECASE)
    p_url   = re.compile(URL_VAL_PDF,   re.IGNORECASE)

    for pageno, page in enumerate(doc, start=1):
        txt = page.get_text() or ""

        # מחיקת קישורי PDF כדי לבטל לחיצות
        try:
            for L in page.get_links():
                try: page.delete_link(L)
                except Exception:
                    if 'xref' in L: page.delete_link(L['xref'])
        except Exception:
            pass

        used_ocr = False
        name_title_hits = 0
        inline_birth_replaced = 0
        nearby_birth_replaced = 0
        added_rects = []

        if not txt.strip():
            # ===== OCR MODE =====
            used_ocr = True
            # 1) שם ככותרת באמצעות OCR
            name_title_hits = _ocr_redact_probable_name_title(page, zoom=3.0)

            # 2) תאריך לידה -> שנה (OCR)
            for rect, year in _ocr_birthdate_to_year_rects(page, zoom=3.0):
                page.add_redact_annot(rect, text=year, fill=(1,1,1))
                added_rects.append(rect)
                inline_birth_replaced += 1

            # 3) אימייל/קישור/טלפון/ת"ז (OCR token windows)
            rects_token_emails = _ocr_find_token_windows(page, re.compile(EMAIL_VAL_TXT, re.IGNORECASE), max_tokens=8,  zoom=3.0)
            rects_token_urls   = _ocr_find_token_windows(page, re.compile(URL_VAL_TXT,   re.IGNORECASE), max_tokens=10, zoom=3.0)
            rects_token_phone  = _ocr_find_token_windows(page, re.compile(PHONE_VAL_PDF), max_tokens=6,   zoom=3.0)
            rects_token_id     = _ocr_find_token_windows(page, re.compile(ID_VAL_PDF),    max_tokens=4,   zoom=3.0)

            rects_all = _merge_overlapping_rects_linewise(rects_token_emails + rects_token_urls + rects_token_phone + rects_token_id)
            for r in rects_all:
                page.add_redact_annot(r, text=REPL, fill=(1,1,1))
                added_rects.append(r)
            _safe_apply_redactions(page, added_rects)

            print(f'PDF page {pageno} (OCR): redacted {len(rects_all)} regions, birth(ocr)={inline_birth_replaced}, name_title={name_title_hits}')
        else:
            # ===== TEXT MODE =====
            # שם ככותרת בטקסט חי
            name_title_hits = _redact_probable_name_title(page)
            # נסה גם לפי גודל פונטים בחלק העליון
            if name_title_hits == 0:
                name_title_hits += _redact_big_font_name_spans(page)
            # ואם עדיין לא – Fallback OCR רק לשם כותרת
            if name_title_hits == 0:
                name_title_hits += _ocr_redact_probable_name_title(page, zoom=3.0)

            # תאריך לידה -> שנה (טקסט מלא ושכנים)
            def _yy2yyyy(yy):
                n=int(yy); return f'20{n:02d}' if n<=24 else f'19{n:02d}'
            date_core = (
                r'(?:'
                r'(?P<yyyy>\d{4})[./\-]\d{1,2}[./\-]\d{1,2}'
                r'|'
                r'\d{1,2}[./\-]\d{1,2}[./\-](?P<yyyy2>\d{4})'
                r'|'
                rf'{ENG_MONTHS}\s+\d{{1,2}},?\s+(?P<yyyy3>\d{{4}})'
                r'|'
                rf'\d{{1,2}}\s+{HEB_MONTHS}\s+(?P<yyyy4>\d{{4}})'
                r'|'
                r'\d{1,2}[./\-]\d{1,2}[./\-](?P<yy>\d{2})'
                r'|'
                rf'{ENG_MONTHS}\s+\d{{1,2}},?\s+(?P<yy2>\d{{2}})'
                r')'
            )
            pat_inline = re.compile(rf'({LABEL_BIRTHDATE}){WSP}[:：\-]?\s*({date_core})', re.IGNORECASE)
            for m in pat_inline.finditer(txt):
                full_date = m.group(2)
                year = (m.group('yyyy') or m.group('yyyy2') or m.group('yyyy3') or
                        m.group('yyyy4') or m.group('yy') or m.group('yy2'))
                if not year: continue
                if len(year) == 2: year = _yy2yyyy(year)
                try:
                    for r in page.search_for(full_date):
                        page.add_redact_annot(r, text=year, fill=(1,1,1)); added_rects.append(r); inline_birth_replaced += 1
                except Exception:
                    pass

            date_re  = re.compile(date_core)
            label_re = re.compile(rf'({LABEL_BIRTHDATE})', re.IGNORECASE)
            for m in date_re.finditer(txt):
                i0, i1 = m.span()
                left  = txt[max(0, i0 - 40): i0]
                right = txt[i1: i1 + 40]
                if not (label_re.search(left) or label_re.search(right)): continue
                full_date = m.group(0)
                year = (m.group('yyyy') or m.group('yyyy2') or m.group('yyyy3') or
                        m.group('yyyy4') or m.group('yy') or m.group('yy2'))
                if not year: continue
                if len(year) == 2: year = _yy2yyyy(year)
                try:
                    for r in page.search_for(full_date):
                        page.add_redact_annot(r, text=year, fill=(1,1,1)); added_rects.append(r); nearby_birth_replaced += 1
                except Exception:
                    pass

            # טקסטים פשוטים (טל/ת"ז/מייל/URL) + tokenwise
            simple_strings = set()
            for pat in (p_phone, p_id, p_mail, p_url):
                for mm in pat.finditer(txt): simple_strings.add(mm.group(0))
            rects_simple = []
            for s in simple_strings:
                try: rects_simple.extend(page.search_for(s))
                except Exception: pass

            rects_name  = _redact_name_values_rects(page)

            # tokenwise (מקרים של שבירת מילים)
            TOKEN_EMAIL_RE = re.compile(EMAIL_VAL_TXT, re.IGNORECASE)
            TOKEN_URL_RE   = re.compile(URL_VAL_TXT,   re.IGNORECASE)
            def _token_rects(page, regex, max_tokens):
                ws = _words(page); n=len(ws); rects=[]
                for i in range(n):
                    for j in range(i, min(n, i+max_tokens)):
                        s0=''.join(ws[k][4] for k in range(i,j+1))
                        s1=' '.join(ws[k][4] for k in range(i,j+1))
                        if regex.search(s0) or regex.search(s1):
                            rects.append(_rect_union(ws,i,j)); break
                return rects
            rects_token_emails = _token_rects(page, TOKEN_EMAIL_RE, 8)
            rects_token_urls   = _token_rects(page, TOKEN_URL_RE, 10)

            rects_all = _merge_overlapping_rects_linewise(rects_simple + rects_name + rects_token_emails + rects_token_urls)
            for r in rects_all:
                page.add_redact_annot(r, text=REPL, fill=(1,1,1)); added_rects.append(r)

            _safe_apply_redactions(page, added_rects)

            print(f'PDF page {pageno}: redacted {len(rects_all)} regions, birth(tokens) inline={inline_birth_replaced}, near={nearby_birth_replaced}, name_title={name_title_hits}')

    doc.save(path_out)
    doc.close()

# ===== עזר: המרות לקבצים ישנים =====
CONVERTABLE_EXTS = {'.doc', '.docm', '.rtf', '.odt'}
SUPPORTED_FINAL  = {'.docx', '.pdf'}

def discover_all_files(root_dir):
    return [str(p) for p in Path(root_dir).rglob('*') if p.is_file()]

def libreoffice_convert(in_path: str, out_fmt: str = 'docx') -> str | None:
    in_path = str(in_path)
    out_dir = str(Path(in_path).parent)
    cmd = f'soffice --headless --convert-to {out_fmt} --outdir {shlex.quote(out_dir)} {shlex.quote(in_path)}'
    try:
        res = subprocess.run(cmd, shell=True, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, timeout=120)
        if res.returncode == 0:
            out_path = str(Path(in_path).with_suffix(f'.{out_fmt}'))
            return out_path if Path(out_path).exists() else None
    except subprocess.TimeoutExpired:
        print(f'WARN: convert timeout: {in_path}')
    except Exception as e:
        print(f'WARN: convert failed: {in_path} | {e}')
    return None

def make_unique_outpath(in_path: str, input_root: str, output_root: str) -> str:
    p_in   = Path(in_path)
    try:
        rel = p_in.relative_to(input_root)
        rel_key = str(rel)
    except Exception:
        rel_key = p_in.name
    hshort = hashlib.sha1(rel_key.encode('utf-8')).hexdigest()[:8]
    return str(Path(output_root) / f"{p_in.stem}__ANON__{hshort}{p_in.suffix.lower()}")

# ===== ריצה =====
all_found = discover_all_files(INPUT_DIR)
print(f'נמצאו בקלט (רקורסיבי): {len(all_found)} קבצים')

# המרות מראש
to_convert = [p for p in all_found if Path(p).suffix.lower() in CONVERTABLE_EXTS]
converted_map = {}
for src in to_convert:
    outp = libreoffice_convert(src, 'docx')
    if outp:
        converted_map[src] = outp
        print(f'Converted -> {outp}')
    else:
        print(f'FAILED convert: {src}')

# רשימת קלט סופית
final_inputs = []
for p in all_found:
    if Path(p).suffix.lower() in SUPPORTED_FINAL:
        final_inputs.append(p)
final_inputs.extend(converted_map.values())
final_inputs = list(dict.fromkeys(final_inputs))

print(f'לקבצים לעיבוד (PDF/DOCX): {len(final_inputs)}')

processed_ok, processed_fail = [], []

for in_path in final_inputs:
    ext = Path(in_path).suffix.lower()
    out_path = make_unique_outpath(in_path, INPUT_DIR, OUTPUT_DIR)
    try:
        if ext == '.docx':
            print('Processing DOCX:', Path(in_path).name, '->', Path(out_path).name)
            anonymize_docx_file(in_path, out_path)
        else:
            print('Processing PDF:', Path(in_path).name, '->', Path(out_path).name)
            anonymize_pdf_file(in_path, out_path)
        processed_ok.append((in_path, out_path))
    except Exception as e:
        processed_fail.append(in_path)
        print(f'FAILED: {Path(in_path).name} | {type(e).__name__}: {e}')

print('\n===== SUMMARY =====')
print(f'סה״כ נמצאו (כל הסוגים): {len(all_found)}')
print(f'ניסיונות המרה: {len(to_convert)} | הומרו בהצלחה: {len(converted_map)}')
print(f'עברו עיבוד (נשמרו ב-{OUTPUT_DIR}): {len(processed_ok)}')
print(f'נכשלו בעיבוד: {len(processed_fail)}')
