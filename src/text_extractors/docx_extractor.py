from typing import Optional
from pathlib import Path
from docx import Document

def extract_text_from_docx(path: str) -> Optional[str]:
    """
    Extract visible text from a DOCX file — paragraphs and tables.
    Returns the plain text or None if failed.
    """
    p = Path(path)
    if not p.exists():
        return None

    try:
        doc = Document(str(p))
    except Exception:
        return None

    parts = []
    # paragraphs
    for para in doc.paragraphs:
        txt = (para.text or "").strip()
        if txt:
            parts.append(txt)

    # tables
    for tbl in doc.tables:
        for row in tbl.rows:
            cells = [(cell.text or "").strip() for cell in row.cells]
            line = " | ".join([c for c in cells if c])
            if line:
                parts.append(line)

    text = "\n".join(parts).strip()
    return text or None
